import hashlib
import io
import re
import zipfile
from pathlib import Path
from xml.etree import ElementTree

from docx import Document
from pypdf import PdfReader

from app.integrations.protocols import DocumentIR, DocumentLocation, DocumentNode

PDF_MIME = "application/pdf"
DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
TXT_MIME = "text/plain"


class LightweightDocumentParser:
    """Parse common office files without Torch, OCR, or local ML models."""

    parser_name = "lightweight"
    parser_version = "1"

    def parse(
        self,
        source: bytes | Path,
        mime_type: str,
        *,
        filename: str | None = None,
    ) -> DocumentIR:
        payload = source.read_bytes() if isinstance(source, Path) else source
        if mime_type == TXT_MIME:
            return self._plain_text(payload)
        if mime_type == PDF_MIME:
            return self._pdf(payload)
        if mime_type == DOCX_MIME:
            return self._docx(payload)
        if mime_type == XLSX_MIME:
            return self._xlsx(payload)
        raise ValueError(f"unsupported parser MIME type: {mime_type}")

    def _plain_text(self, payload: bytes) -> DocumentIR:
        text = payload.decode("utf-8", errors="strict")
        nodes = tuple(
            _node(
                value,
                "paragraph",
                {
                    "kind": "plain_text",
                    "start_offset": match.start(),
                    "end_offset": match.end(),
                },
            )
            for match in re.finditer(r"[^\r\n]+", text)
            if (value := match.group(0).strip())
        )
        return self._result(nodes)

    def _pdf(self, payload: bytes) -> DocumentIR:
        nodes: list[DocumentNode] = []
        warnings: list[str] = []
        reader = PdfReader(io.BytesIO(payload))
        for page_number, page in enumerate(reader.pages, start=1):
            text = (page.extract_text() or "").strip()
            if text:
                nodes.append(_node(text, "paragraph", {"kind": "pdf_page", "page": page_number}))
            else:
                warnings.append(f"page {page_number} has no extractable text; OCR is not enabled")
        return self._result(tuple(nodes), tuple(warnings))

    def _docx(self, payload: bytes) -> DocumentIR:
        document = Document(io.BytesIO(payload))
        nodes: list[DocumentNode] = []
        for index, paragraph in enumerate(document.paragraphs):
            text = paragraph.text.strip()
            if not text:
                continue
            style = (paragraph.style.name if paragraph.style else "").lower()
            node_type = "heading" if "heading" in style or "标题" in style else "paragraph"
            nodes.append(
                _node(
                    text,
                    node_type,
                    {"kind": "docx_paragraph", "paragraph_index": index},
                )
            )
        for table_index, table in enumerate(document.tables):
            rows = [
                [cell.text.strip().replace("\n", " ") for cell in row.cells] for row in table.rows
            ]
            text = _markdown_table(rows)
            if text:
                nodes.append(
                    _node(
                        text,
                        "table",
                        {"kind": "docx_paragraph", "table_index": table_index},
                    )
                )
        return self._result(tuple(nodes))

    def _xlsx(self, payload: bytes) -> DocumentIR:
        nodes: list[DocumentNode] = []
        with zipfile.ZipFile(io.BytesIO(payload)) as archive:
            shared = _xlsx_shared_strings(archive)
            sheet_names = sorted(
                name
                for name in archive.namelist()
                if re.fullmatch(r"xl/worksheets/sheet\d+\.xml", name)
            )
            for sheet_number, name in enumerate(sheet_names, start=1):
                root = ElementTree.fromstring(archive.read(name))
                namespace = _xml_namespace(root.tag)
                for row in root.findall(f".//{{{namespace}}}row"):
                    values: list[str] = []
                    references: list[str] = []
                    for cell in row.findall(f"{{{namespace}}}c"):
                        references.append(cell.attrib.get("r", ""))
                        values.append(_xlsx_cell_value(cell, namespace, shared))
                    if not any(values):
                        continue
                    start = references[0] if references else f"A{sheet_number}"
                    end = references[-1] if references else start
                    nodes.append(
                        _node(
                            " | ".join(values),
                            "table",
                            {
                                "kind": "xlsx_cell",
                                "sheet": f"Sheet{sheet_number}",
                                "cell_range": f"{start}:{end}" if start != end else start,
                            },
                        )
                    )
        return self._result(tuple(nodes))

    def _result(
        self, nodes: tuple[DocumentNode, ...], warnings: tuple[str, ...] = ()
    ) -> DocumentIR:
        return DocumentIR(
            nodes=nodes,
            parser_name=self.parser_name,
            parser_version=self.parser_version,
            warnings=warnings,
        )


def _node(text: str, node_type: str, location: DocumentLocation) -> DocumentNode:
    located = {
        "schema_version": "document-location-v1",
        "quote_hash": hashlib.sha256(text.encode("utf-8")).hexdigest(),
        **location,
    }
    return DocumentNode(text=text, node_type=node_type, location=located)  # type: ignore[arg-type]


def _markdown_table(rows: list[list[str]]) -> str:
    if not rows:
        return ""
    width = max(len(row) for row in rows)
    normalized = [row + [""] * (width - len(row)) for row in rows]
    head = normalized[0]
    body = normalized[1:]
    rendered = ["| " + " | ".join(head) + " |", "| " + " | ".join(["---"] * width) + " |"]
    rendered.extend("| " + " | ".join(row) + " |" for row in body)
    return "\n".join(rendered)


def _xml_namespace(tag: str) -> str:
    return tag[1:].split("}", 1)[0] if tag.startswith("{") else ""


def _xlsx_shared_strings(archive: zipfile.ZipFile) -> list[str]:
    try:
        root = ElementTree.fromstring(archive.read("xl/sharedStrings.xml"))
    except KeyError:
        return []
    namespace = _xml_namespace(root.tag)
    return [
        "".join(node.text or "" for node in item.findall(f".//{{{namespace}}}t"))
        for item in root.findall(f"{{{namespace}}}si")
    ]


def _xlsx_cell_value(cell, namespace: str, shared: list[str]) -> str:
    kind = cell.attrib.get("t")
    if kind == "inlineStr":
        return "".join(node.text or "" for node in cell.findall(f".//{{{namespace}}}t"))
    value = cell.find(f"{{{namespace}}}v")
    raw = value.text if value is not None and value.text is not None else ""
    if kind == "s" and raw.isdigit() and int(raw) < len(shared):
        return shared[int(raw)]
    return raw
