import asyncio
import hashlib
import io
import uuid
import zipfile
from types import SimpleNamespace

import pytest
from docx import Document

from app.core.errors import AppError, ErrorCode
from app.db.models import TenderParseStatus
from app.integrations.docling_adapter import DoclingAdapter
from app.integrations.lightweight_document_parser import LightweightDocumentParser
from app.integrations.protocols import DocumentIR, DocumentNode
from app.services.tender_parse_worker import (
    MAX_TENDER_FILE_BYTES,
    TenderParseWorker,
    _configured_parser,
    read_and_validate_tender_source,
)


def pdf_location(text: str) -> dict:
    return {
        "kind": "pdf_page",
        "schema_version": "document-location-v1",
        "quote_hash": hashlib.sha256(text.encode()).hexdigest(),
        "page": 1,
        "bounding_box": [1.0, 2.0, 30.0, 40.0],
    }


class FakeSession:
    async def commit(self) -> None:
        pass

    async def rollback(self) -> None:
        pass


class FakeEntityRepository:
    def __init__(self, entity) -> None:
        self.entity = entity

    async def get(self, entity_id):
        return self.entity if self.entity.id == entity_id else None

    async def get_for_update(self, entity_id):
        return await self.get(entity_id)


class FakeParseRepository:
    def __init__(self) -> None:
        self.entity = None

    async def next_version(self, tender_id) -> int:
        return 1

    async def create(self, **values):
        self.entity = SimpleNamespace(id=uuid.uuid4(), **values)
        return self.entity

    async def get(self, entity_id):
        return self.entity if self.entity and self.entity.id == entity_id else None


class FakeRequirementRepository:
    def __init__(self) -> None:
        self.created: list[dict] = []

    async def soft_delete_for_tender(self, tender_id) -> None:
        pass

    async def next_sequence(self, tender_id) -> int:
        return 1

    async def create(self, **values):
        self.created.append(values)


class MockParser:
    parser_name = "mock-docling"
    parser_version = "1.0"

    def parse(self, source, mime_type, *, filename=None) -> DocumentIR:
        text = "供应商必须提供三年成功案例。"
        return DocumentIR(
            nodes=(DocumentNode(text, "paragraph", pdf_location(text)),),
            parser_name=self.parser_name,
            parser_version=self.parser_version,
        )


class SlowParser(MockParser):
    async def parse(self, source, mime_type, *, filename=None) -> DocumentIR:
        await asyncio.sleep(400)
        raise AssertionError("unreachable")


def configured_worker(parser, *, timeout_seconds=1.0):
    tender = SimpleNamespace(id=uuid.uuid4())
    artifact = SimpleNamespace(id=uuid.uuid4())
    worker = TenderParseWorker(
        FakeSession(),
        uuid.uuid4(),
        parser=parser,
        timeout_seconds=timeout_seconds,
    )
    worker.tenders = FakeEntityRepository(tender)
    worker.artifacts = FakeEntityRepository(artifact)
    worker.parse_versions = FakeParseRepository()
    worker.requirements = FakeRequirementRepository()
    return worker, tender, artifact


@pytest.mark.asyncio
async def test_parse_state_machine_creates_version_and_located_requirements() -> None:
    worker, tender, artifact = configured_worker(MockParser())

    version = await worker.parse(
        tender_id=tender.id,
        raw_artifact_id=artifact.id,
        source=b"%PDF-1.7 minimal",
        mime_type="application/pdf",
        filename="tender.pdf",
    )

    assert version.status == TenderParseStatus.PARSED
    assert version.document_ir["nodes"][0]["location"]["page"] == 1
    assert version.resource_usage["requirement_count"] == 1
    requirement = worker.requirements.created[0]
    assert requirement["parse_version_id"] == version.id
    assert requirement["source_location"]["bounding_box"] == [1.0, 2.0, 30.0, 40.0]


@pytest.mark.asyncio
async def test_oversized_file_fails_before_parser_is_called() -> None:
    class UnexpectedParser(MockParser):
        def parse(self, source, mime_type, *, filename=None) -> DocumentIR:
            raise AssertionError("parser must not be called")

    worker, tender, artifact = configured_worker(UnexpectedParser())

    with pytest.raises(AppError) as caught:
        await worker.parse(
            tender_id=tender.id,
            raw_artifact_id=artifact.id,
            source=b"x" * (MAX_TENDER_FILE_BYTES + 1),
            mime_type="application/pdf",
            filename="large.pdf",
        )

    assert caught.value.code == ErrorCode.TENDER_FILE_TOO_LARGE
    assert worker.parse_versions.entity.status == TenderParseStatus.FAILED
    assert worker.parse_versions.entity.error_code == "FILE_TOO_LARGE"


@pytest.mark.asyncio
async def test_parser_timeout_is_bounded_and_persisted() -> None:
    worker, tender, artifact = configured_worker(SlowParser(), timeout_seconds=0.01)

    with pytest.raises(AppError) as caught:
        await worker.parse(
            tender_id=tender.id,
            raw_artifact_id=artifact.id,
            source=b"%PDF-1.7 minimal",
            mime_type="application/pdf",
            filename="slow.pdf",
        )

    assert caught.value.code == ErrorCode.TENDER_PARSE_TIMEOUT
    assert worker.parse_versions.entity.status == TenderParseStatus.FAILED
    assert worker.parse_versions.entity.error_code == "TIMEOUT"


@pytest.mark.asyncio
async def test_mime_probe_rejects_executable_disguised_as_pdf() -> None:
    worker, tender, artifact = configured_worker(MockParser())

    with pytest.raises(AppError) as caught:
        await worker.parse(
            tender_id=tender.id,
            raw_artifact_id=artifact.id,
            source=b"MZ executable",
            mime_type="application/pdf",
            filename="disguised.pdf",
        )

    assert caught.value.code == ErrorCode.TENDER_FILE_UNSAFE
    assert worker.parse_versions.entity.status == TenderParseStatus.FAILED


def test_docling_maps_pdf_table_to_markdown_and_bounding_box() -> None:
    class Table:
        label = "table"
        prov = [{"page_no": 2, "bbox": {"l": 1, "t": 2, "r": 10, "b": 20}}]

        def export_to_markdown(self) -> str:
            return "| A | B |\n|---|---|\n| 1 | 2 |"

    document = {"items": [Table()]}
    result = DoclingAdapter(converter=object()).map_document(document, "application/pdf")

    node = result.nodes[0]
    assert node.node_type == "table"
    assert node.text.startswith("| A | B |")
    assert node.location["page"] == 2
    assert node.location["bounding_box"] == [1.0, 2.0, 10.0, 20.0]
    assert node.location["content_kind"] == "table"


def test_utf8_txt_is_parsed_with_verifiable_offsets() -> None:
    payload = "供应商必须提供三年案例。\n系统应支持不少于1000并发。".encode()

    validated = read_and_validate_tender_source(payload, "text/plain", "requirements.txt")
    result = DoclingAdapter().parse(validated, "text/plain", filename="requirements.txt")

    assert len(result.nodes) == 2
    assert result.nodes[0].location["kind"] == "plain_text"
    assert result.nodes[0].location["start_offset"] == 0
    assert (
        result.nodes[0].location["quote_hash"]
        == hashlib.sha256(result.nodes[0].text.encode()).hexdigest()
    )


def test_default_parser_does_not_require_docling_or_torch() -> None:
    parser = _configured_parser()

    assert isinstance(parser, LightweightDocumentParser)
    assert parser.parser_name == "lightweight"


def test_lightweight_parser_reads_docx_without_local_ml_models() -> None:
    document = Document()
    document.add_heading("采购要求", level=1)
    document.add_paragraph("供应商必须提供三年成功案例。")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "指标"
    table.cell(0, 1).text = "要求"
    table.cell(1, 0).text = "并发"
    table.cell(1, 1).text = "1000"
    stream = io.BytesIO()
    document.save(stream)

    result = LightweightDocumentParser().parse(
        stream.getvalue(),
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        filename="requirements.docx",
    )

    assert any(node.node_type == "heading" and node.text == "采购要求" for node in result.nodes)
    assert any(node.node_type == "table" and "1000" in node.text for node in result.nodes)


def test_lightweight_parser_reads_xlsx_xml_without_openpyxl() -> None:
    stream = io.BytesIO()
    with zipfile.ZipFile(stream, "w") as archive:
        archive.writestr(
            "xl/sharedStrings.xml",
            '<?xml version="1.0"?><sst xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"><si><t>必须支持国产化</t></si></sst>',
        )
        archive.writestr(
            "xl/worksheets/sheet1.xml",
            '<?xml version="1.0"?>'
            '<worksheet xmlns="http://schemas.openxmlformats.org/'
            'spreadsheetml/2006/main"><sheetData><row r="1">'
            '<c r="A1" t="s"><v>0</v></c><c r="B1"><v>1</v></c>'
            "</row></sheetData></worksheet>",
        )

    result = LightweightDocumentParser().parse(
        stream.getvalue(),
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        filename="requirements.xlsx",
    )

    assert result.nodes[0].text == "必须支持国产化 | 1"
    assert result.nodes[0].location["cell_range"] == "A1:B1"


def test_txt_rejects_non_utf8_content() -> None:
    with pytest.raises(AppError) as caught:
        read_and_validate_tender_source(b"\xff\xfe", "text/plain", "requirements.txt")

    assert caught.value.code == ErrorCode.TENDER_FILE_UNSAFE


def test_office_archive_rejects_unsafe_member_path() -> None:
    stream = io.BytesIO()
    with zipfile.ZipFile(stream, "w") as archive:
        archive.writestr("[Content_Types].xml", "<Types/>")
        archive.writestr("word/document.xml", "<document/>")
        archive.writestr("../outside", "unsafe")

    with pytest.raises(AppError) as caught:
        read_and_validate_tender_source(
            stream.getvalue(),
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            "requirements.docx",
        )

    assert caught.value.code == ErrorCode.TENDER_FILE_UNSAFE
