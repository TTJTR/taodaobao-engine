"""Run the V2 tender parser and response Trust Gate against an isolated workspace."""

from __future__ import annotations

import asyncio
import hashlib
import importlib.util
import sys
import tempfile
import uuid
from datetime import UTC, datetime
from pathlib import Path

from sqlalchemy import inspect as sa_inspect
from sqlalchemy import select, text, update
from sqlalchemy.ext.asyncio import create_async_engine

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from app.core.config import settings  # noqa: E402
from app.db.database import create_db_engine, create_session_factory  # noqa: E402
from app.db.models import (  # noqa: E402
    Base,
    Capability,
    RawArtifact,
    RawArtifactKind,
    RawArtifactStatus,
    ReviewStatus,
    Source,
    SourceFreshness,
    SourcePurpose,
    SourceStatus,
    SourceType,
    TenderDocument,
    TenderParseStatus,
    TenderRequirement,
    User,
)
from app.integrations.docling_adapter import DoclingAdapter  # noqa: E402
from app.integrations.protocols import DocumentNode  # noqa: E402
from app.services.tender_parse_worker import TenderParseWorker  # noqa: E402
from app.services.tender_service import TenderService  # noqa: E402

MIME_DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
REQUIREMENT_SUPPORTED = "要求系统必须支持 10 万并发处理。"
REQUIREMENT_UNSUPPORTED = "要求系统必须提供宇宙飞船对接接口。"


def _node(value: str, paragraph: int, node_type: str, *, row: int) -> DocumentNode:
    return DocumentNode(
        text=value,
        node_type=node_type,
        location={
            "schema_version": "document-location-v1",
            "kind": "docx_paragraph",
            "section_path": ["技术要求", "并发与特殊接口"],
            "paragraph_index": paragraph,
            "table_index": 0,
            "row_index": row,
            "column_index": 1,
            "quote_hash": hashlib.sha256(value.encode()).hexdigest(),
        },
    )


def generate_docx(path: Path) -> None:
    try:
        from docx import Document
    except ImportError as exc:
        raise RuntimeError("缺少 python-docx，请先安装项目依赖") from exc
    document = Document()
    document.add_heading("V2 招标系统 E2E 验收文件", level=1)
    document.add_paragraph("以下技术要求用于验证内部证据隔离和缺口拦截。")
    table = document.add_table(rows=3, cols=2)
    table.style = "Table Grid"
    table.cell(0, 0).text = "编号"
    table.cell(0, 1).text = "技术要求"
    table.cell(1, 0).text = "REQ-001"
    table.cell(1, 1).text = REQUIREMENT_SUPPORTED
    table.cell(2, 0).text = "REQ-002"
    table.cell(2, 1).text = REQUIREMENT_UNSUPPORTED
    document.save(path)


async def main() -> int:
    if not importlib.util.find_spec("docling"):
        print("[SKIPPED] Docling is not installed; tender parsing was not simulated.")
        return 0
    if not settings.database_url:
        print("[FAIL] 未配置 APP_DATABASE_URL，无法连接 PostgreSQL。")
        return 2
    workspace_id = uuid.uuid4()
    schema_name = f"tender_e2e_{workspace_id.hex}"
    admin_engine = create_db_engine(settings.database_url)
    engine = None
    factory = None
    temp_dir = tempfile.TemporaryDirectory(prefix="tender-e2e-")
    docx_path = Path(temp_dir.name) / "synthetic-tender.docx"
    print(f"[SETUP] E2E workspace: {workspace_id}")
    try:
        async with admin_engine.begin() as connection:
            await connection.execute(text("SELECT 1"))
            await connection.execute(text(f'CREATE SCHEMA "{schema_name}"'))
        engine = create_async_engine(
            settings.database_url,
            pool_pre_ping=True,
            connect_args={"server_settings": {"search_path": f"{schema_name}, public"}},
        )
        factory = create_session_factory(engine)
        async with engine.begin() as connection:
            await connection.run_sync(Base.metadata.create_all)
        async with factory() as session:
            print("[OK] PostgreSQL 连接正常")
            print(f"[OK] 已创建隔离验收 Schema: {schema_name}")
            generate_docx(docx_path)
            payload = docx_path.read_bytes()
            print(f"[OK] 已动态生成 DOCX: {docx_path.name} ({len(payload)} bytes)")

            user = User(
                workspace_id=workspace_id,
                feishu_user_id=f"tender-e2e-{workspace_id}",
                name="Tender E2E Reviewer",
            )
            session.add(user)
            await session.flush()
            source = Source(
                workspace_id=workspace_id,
                imported_by_id=user.id,
                type=SourceType.PASTED_TEXT,
                purpose=SourcePurpose.CAPABILITY,
                title="已验证高并发能力说明",
                content="系统支持 10 万并发处理，并经过内部性能验证。",
                source_url=f"e2e://capability/{workspace_id}",
                status=SourceStatus.COMPLETED,
                freshness_status=SourceFreshness.CURRENT,
                content_fingerprint=hashlib.sha256(b"100k-concurrency").hexdigest(),
                content_version=1,
                permission_checked_at=datetime.now(UTC),
            )
            session.add(source)
            await session.flush()
            capability = Capability(
                workspace_id=workspace_id,
                source_id=source.id,
                data={"name": "10 万并发处理能力", "description": "支持 10 万并发处理"},
                review_status=ReviewStatus.VERIFIED,
                reviewed_by_id=user.id,
                reviewed_at=datetime.now(UTC),
                source_version_at_review=1,
                embedding=[0.0] * 1024,
                embedding_text="支持 10 万并发处理",
                embedding_version="e2e-zero-vector-v1",
                embedding_fingerprint=hashlib.sha256(b"e2e-embedding").hexdigest(),
                embedding_ready=True,
            )
            session.add(capability)
            tender = TenderDocument(
                workspace_id=workspace_id,
                created_by_id=user.id,
                title="V2 招标 E2E 验收",
                source_filename=docx_path.name,
                source_mime_type=MIME_DOCX,
                source_fingerprint=hashlib.sha256(payload).hexdigest(),
                content_text=f"{REQUIREMENT_SUPPORTED}\n{REQUIREMENT_UNSUPPORTED}",
                status="ready",
            )
            session.add(tender)
            await session.flush()
            artifact = RawArtifact(
                workspace_id=workspace_id,
                artifact_key=hashlib.sha256(payload).hexdigest(),
                kind=RawArtifactKind.TENDER_FILE,
                status=RawArtifactStatus.CAPTURED,
                provider="e2e-script",
                source_filename=docx_path.name,
                mime_type=MIME_DOCX,
                content_sha256=hashlib.sha256(payload).hexdigest(),
                byte_size=len(payload),
                storage_uri=docx_path.resolve().as_uri(),
                captured_at=datetime.now(UTC),
                security_report={"generated_by": "verify_tender_e2e.py"},
                metadata_snapshot={},
            )
            session.add(artifact)
            await session.flush()
            tender.raw_artifact_id = artifact.id
            await session.commit()
            print(f"[STEP 1] RawArtifact: {artifact.id}")

            parser = DoclingAdapter()
            print(f"[STEP 2] Parser: {parser.parser_name} {parser.parser_version}")
            parsed = await TenderParseWorker(
                session, workspace_id, parser=parser
            ).parse(
                tender_id=tender.id,
                raw_artifact_id=artifact.id,
                source=docx_path,
                mime_type=MIME_DOCX,
                filename=docx_path.name,
            )
            if parsed.status != TenderParseStatus.PARSED:
                raise RuntimeError(f"解析未完成: {parsed.status}")
            print(f"[OK] TenderParseVersion: {parsed.id} status={parsed.status.value}")

            service = TenderService(session, workspace_id, user.id)
            requirements = await service.breakdown_requirements(tender.id)
            print(f"[STEP 3] 拆解出 {len(requirements)} 条 TenderRequirement")
            matrix = await service.create_matrix(
                tender.id,
                experience_ids=[],
                capability_ids=[capability.id],
                intelligence_snapshot_id=None,
            )
            _, items = await service.get_matrix(matrix.id)
            requirement_rows = {
                row.id: row
                for row in await session.scalars(
                    select(TenderRequirement).where(
                        TenderRequirement.workspace_id == workspace_id,
                        TenderRequirement.tender_id == tender.id,
                        TenderRequirement.is_deleted.is_(False),
                    )
                )
            }
            rows = [
                {
                    "requirement": requirement_rows[item.requirement_id].requirement_text,
                    "location": _location(requirement_rows[item.requirement_id].source_location),
                    "evidence": ", ".join(
                        str(link.get("asset_id")) for link in item.internal_cap_links
                    ) or "-",
                    "draft": item.ai_draft,
                    "risks": ", ".join(item.risk_flags) or "-",
                }
                for item in items
            ]
            print("[STEP 4] Response Matrix Trust Gate")
            _print_table(rows)
            spaceship = next(row for row in rows if "宇宙飞船" in row["requirement"])
            supported = next(row for row in rows if "并发" in row["requirement"])
            if spaceship["risks"] != "missing_evidence" or supported["evidence"] == "-":
                raise RuntimeError("Trust Gate 验收失败：证据隔离结果不符合预期")
            print("[PASS] 宇宙飞船要求被标记 missing_evidence，未生成企业能力承诺。")
            return 0
    except Exception as exc:
        print(f"[FAIL] {type(exc).__name__}: {exc}")
        return 1
    finally:
        try:
            if factory is None:
                raise RuntimeError("隔离数据库尚未初始化")
            async with factory() as cleanup_session:
                count = await _soft_delete_workspace(cleanup_session, workspace_id)
                print(f"[CLEANUP] 已软删除测试 Workspace 下 {count} 条记录")
        except Exception as cleanup_error:
            print(f"[WARN] 清理失败: {cleanup_error}")
        if engine is not None:
            await engine.dispose()
        try:
            async with admin_engine.begin() as connection:
                await connection.execute(text(f'DROP SCHEMA IF EXISTS "{schema_name}" CASCADE'))
            print(f"[CLEANUP] 已删除隔离验收 Schema: {schema_name}")
        except Exception as schema_error:
            print(f"[WARN] 删除隔离 Schema 失败: {schema_error}")
        await admin_engine.dispose()
        temp_dir.cleanup()


async def _soft_delete_workspace(session, workspace_id: uuid.UUID) -> int:
    affected = 0
    existing_tables = set(
        await session.run_sync(lambda sync: sa_inspect(sync.connection()).get_table_names())
    )
    for mapper in reversed(list(Base.registry.mappers)):
        model = mapper.class_
        if not hasattr(model, "workspace_id") or not hasattr(model, "is_deleted"):
            continue
        if model.__table__.name not in existing_tables:
            continue
        result = await session.execute(
            update(model)
            .where(model.workspace_id == workspace_id, model.is_deleted.is_(False))
            .values(is_deleted=True)
        )
        affected += int(result.rowcount or 0)
    await session.commit()
    return affected


def _location(location: dict) -> str:
    return (
        f"{location.get('kind')} paragraph={location.get('paragraph_index')} "
        f"table={location.get('table_index')} row={location.get('row_index')}"
    )


def _print_table(rows: list[dict[str, str]]) -> None:
    try:
        from rich.console import Console
        from rich.table import Table
    except ImportError:
        _print_plain_table(rows)
        return
    table = Table(title="V2 Tender Response Matrix")
    for title in ("招标要求原文", "原文坐标", "内部证据", "AI 草稿", "风险与缺口"):
        table.add_column(title, overflow="fold")
    for row in rows:
        table.add_row(
            row["requirement"], row["location"], row["evidence"], row["draft"], row["risks"]
        )
    Console().print(table)


def _print_plain_table(rows: list[dict[str, str]]) -> None:
    headers = ("招标要求原文", "原文坐标", "内部证据", "AI 草稿", "风险与缺口")
    keys = ("requirement", "location", "evidence", "draft", "risks")
    widths = (34, 47, 36, 44, 20)
    line = "+" + "+".join("-" * (width + 2) for width in widths) + "+"
    print(line)
    header_cells = (
        f" {title:<{width}} " for title, width in zip(headers, widths, strict=True)
    )
    print("|" + "|".join(header_cells) + "|")
    print(line)
    for row in rows:
        values = [str(row[key]) for key in keys]
        cells = (
            f" {value[:width]:<{width}} "
            for value, width in zip(values, widths, strict=True)
        )
        print("|" + "|".join(cells) + "|")
    print(line)


if __name__ == "__main__":
    raise SystemExit(asyncio.run(main()))
